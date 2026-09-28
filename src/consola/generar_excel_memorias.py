# -*- coding: utf-8 -*-
"""
Script para generar el archivo Excel con las Memorias de Cálculo Completas
del Taller de Perceptrón Simple (Casos 1 y 2).
Genera: docs/Memorias_Calculo_Perceptron_Corte2.xlsx
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def crear_memorias_excel():
    """
    Genera programaticamente el libro de calculo Excel con tres hojas formales:
    1. Resumen Ejecutivo (metricas consolidadas y separabilidad lineal).
    2. Caso 1 Paloma (traza numerica paso a paso de las 3 epocas).
    3. Caso 2 Diagnostico (traza numerica paso a paso de las 3 epocas).
    Guarda el archivo en docs/Memorias_Calculo_Perceptron_Corte2.xlsx.
    """
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto
    wb.remove(wb.active)

    # Estilos comunes
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    sub_font = Font(name="Calibri", size=11, bold=True, color="000000")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # ==========================================
    # HOJA 1: RESUMEN GENERAL
    # ==========================================
    ws0 = wb.create_sheet(title="Resumen Ejecutivo")
    ws0.views.sheetView[0].showGridLines = True

    ws0["A1"] = "UNIVERSIDAD DE CARTAGENA - FACULTAD DE INGENIERÍA"
    ws0["A1"].font = title_font
    ws0["A2"] = "INTELIGENCIA ARTIFICIAL - TALLER DE PERCEPTRÓN SIMPLE (CORTE 2)"
    ws0["A2"].font = sub_font
    ws0["A3"] = "Estudiante: Dago David Palmera Navarro (Cód. 0222321003)"
    ws0["A3"].font = regular_font

    headers_resumen = ["Caso de Estudio", "W(0) Inicial", "b(0) Inicial", "Épocas", "W* Final", "b* Final", "Error Final", "Separabilidad Lineal"]
    for col_num, h in enumerate(headers_resumen, 1):
        c = ws0.cell(row=5, column=col_num)
        c.value = h
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_center

    data_resumen = [
        ["Caso 1: Condicionamiento de la Paloma", "[0.3, -0.9, -0.4]", "0.2", "3", "[2.3, 5.1, -2.4]", "6.2", "0 (0%)", "Linealmente Separable (100%)"],
        ["Caso 2: Diagnóstico Médico (Principal)", "[-0.4, -0.7, 0.3]", "0.1", "3", "[3.6, 3.3, 4.3]", "0.1", "0 (0%)", "Linealmente Separable (100%)"],
        ["Caso 2: Diagnóstico (Optimizado 2 Épocas)", "[-0.9, -0.9, -0.9]", "0.9", "2", "[1.1, 1.1, 1.1]", "-1.1", "0 (0%)", "Linealmente Separable (100%)"],
        ["Caso 2: Diagnóstico (Solución Directa 1 Época)", "[1.0, 1.0, 1.0]", "0.0", "1", "[1.0, 1.0, 1.0]", "0.0", "0 (0%)", "Linealmente Separable (100%)"]
    ]

    for row_idx, row_data in enumerate(data_resumen, 6):
        for col_idx, val in enumerate(row_data, 1):
            c = ws0.cell(row=row_idx, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            c.alignment = align_left if col_idx == 1 else align_center

    # ==========================================
    # HOJA 2: CASO 1 - LA PALOMA
    # ==========================================
    ws1 = wb.create_sheet(title="Caso 1 - Paloma")
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "CASO DE ESTUDIO 1: CONDICIONAMIENTO DE LA PALOMA"
    ws1["A1"].font = title_font
    ws1["A2"] = "Ecuación del Hiperplano: 2.3*x1 + 5.1*x2 - 2.4*x3 + 6.2 = 0"
    ws1["A2"].font = sub_font

    headers_iter = ["Época", "Iteración", "Patrón", "x1", "x2", "x3", "W1_antes", "W2_antes", "W3_antes", "b_antes", "a (Suma Neta)", "y (Salida)", "yd (Target)", "e (Error)", "W1_desp", "W2_desp", "W3_desp", "b_desp", "Actualización"]
    for col_num, h in enumerate(headers_iter, 1):
        c = ws1.cell(row=4, column=col_num, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_center

    patterns1 = [
        ([-1, -1, -1], 1),
        ([-1, -1,  1], -1),
        ([-1,  1, -1], 1),
        ([-1,  1,  1], 1),
        ([ 1, -1, -1], 1),
        ([ 1, -1,  1], 1),
        ([ 1,  1, -1], 1),
        ([ 1,  1,  1], 1),
    ]

    W = [0.3, -0.9, -0.4]
    b = 0.2
    row_curr = 5

    for ep in range(1, 4):
        for k, (X, yd) in enumerate(patterns1):
            W_before = W[:]
            b_before = b
            a = round(W[0]*X[0] + W[1]*X[1] + W[2]*X[2] + b, 4)
            y = 1 if a >= 0 else -1
            e = yd - y
            upd = "NO"
            if e != 0:
                W = [round(W[0] + e*X[0], 4), round(W[1] + e*X[1], 4), round(W[2] + e*X[2], 4)]
                b = round(b + e, 4)
                upd = f"SÍ (ΔW={e*X[0]},{e*X[1]},{e*X[2]})"

            row_vals = [ep, k+1, f"P{k+1}", X[0], X[1], X[2], W_before[0], W_before[1], W_before[2], b_before, a, y, yd, e, W[0], W[1], W[2], b, upd]
            for col_idx, val in enumerate(row_vals, 1):
                c = ws1.cell(row=row_curr, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                c.alignment = align_center
                if upd != "NO" and col_idx == 19:
                    c.font = Font(name="Calibri", size=11, bold=True, color="C00000")
            row_curr += 1

    # ==========================================
    # HOJA 3: CASO 2 - DIAGNÓSTICO MÉDICO
    # ==========================================
    ws2 = wb.create_sheet(title="Caso 2 - Diagnóstico")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "CASO DE ESTUDIO 2: DIAGNÓSTICO MÉDICO POR SÍNTOMAS"
    ws2["A1"].font = title_font
    ws2["A2"] = "Ecuación del Hiperplano: 3.6*x1 + 3.3*x2 + 4.3*x3 + 0.1 = 0"
    ws2["A2"].font = sub_font

    for col_num, h in enumerate(headers_iter, 1):
        c = ws2.cell(row=4, column=col_num, value=h)
        c.font = header_font
        c.fill = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
        c.alignment = align_center

    patterns2 = [
        ([-1, -1, -1], -1),
        ([-1, -1,  1], -1),
        ([-1,  1, -1], -1),
        ([-1,  1,  1],  1),
        ([ 1, -1, -1], -1),
        ([ 1, -1,  1],  1),
        ([ 1,  1, -1],  1),
        ([ 1,  1,  1],  1),
    ]

    W = [-0.4, -0.7, 0.3]
    b = 0.1
    row_curr = 5

    for ep in range(1, 4):
        for k, (X, yd) in enumerate(patterns2):
            W_before = W[:]
            b_before = b
            a = round(W[0]*X[0] + W[1]*X[1] + W[2]*X[2] + b, 4)
            y = 1 if a >= 0 else -1
            e = yd - y
            upd = "NO"
            if e != 0:
                W = [round(W[0] + e*X[0], 4), round(W[1] + e*X[1], 4), round(W[2] + e*X[2], 4)]
                b = round(b + e, 4)
                upd = f"SÍ (ΔW={e*X[0]},{e*X[1]},{e*X[2]})"

            row_vals = [ep, k+1, f"P{k+1}", X[0], X[1], X[2], W_before[0], W_before[1], W_before[2], b_before, a, y, yd, e, W[0], W[1], W[2], b, upd]
            for col_idx, val in enumerate(row_vals, 1):
                c = ws2.cell(row=row_curr, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                c.alignment = align_center
                if upd != "NO" and col_idx == 19:
                    c.font = Font(name="Calibri", size=11, bold=True, color="C00000")
            row_curr += 1

    # Ajustar ancho de columnas para todas las hojas
    for ws in [ws0, ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    # Guardar en docs
    ruta_salida = os.path.join(os.path.dirname(__file__), "..", "..", "docs", "Memorias_Calculo_Perceptron_Corte2.xlsx")
    wb.save(ruta_salida)
    print(f"Archivo Excel generado con exito en: {ruta_salida}")

if __name__ == "__main__":
    crear_memorias_excel()
